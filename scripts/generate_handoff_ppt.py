#!/usr/bin/env python3
"""Generate PageIndex handoff PowerPoint deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT = PROJECT_ROOT / "docs" / "PageIndex_Handoff.pptx"

# VTTI-ish teal + slate palette
TEAL = RGBColor(0x0D, 0x94, 0x88)
DARK = RGBColor(0x1E, 0x29, 0x3B)
SLATE = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, TEAL)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(8.8), Inches(1.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    box2 = slide.shapes.add_textbox(Inches(0.6), Inches(3.5), Inches(8.8), Inches(1.5))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(20)
    p2.font.color.rgb = LIGHT
    p2.alignment = PP_ALIGN.LEFT


def add_section_slide(prs: Presentation, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(3.0), Inches(8.8), Inches(1))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE


def add_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    *,
    subtitle: str | None = None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    # title bar
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.9))  # rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.7))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    body_top = 1.15 if not subtitle else 1.45
    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.55), Inches(1.0), Inches(9), Inches(0.4))
        sp = sbox.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.italic = True
        sp.font.color.rgb = SLATE

    bbox = slide.shapes.add_textbox(Inches(0.55), Inches(body_top), Inches(9), Inches(5.8))
    tf = bbox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18 if len(bullet) < 90 else 16)
        p.font.color.rgb = DARK
        p.space_after = Pt(10)


def add_two_column_slide(prs: Presentation, title: str, left: list[str], right: list[str], left_hdr: str, right_hdr: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.7))
    tbox.text_frame.paragraphs[0].text = title
    tbox.text_frame.paragraphs[0].font.size = Pt(28)
    tbox.text_frame.paragraphs[0].font.bold = True
    tbox.text_frame.paragraphs[0].font.color.rgb = WHITE

    for col, hdr, items, x in [(left, left_hdr, left, 0.4), (right, right_hdr, right, 5.1)]:
        hbox = slide.shapes.add_textbox(Inches(x), Inches(1.1), Inches(4.5), Inches(0.4))
        hp = hbox.text_frame.paragraphs[0]
        hp.text = hdr
        hp.font.bold = True
        hp.font.size = Pt(20)
        hp.font.color.rgb = TEAL
        bbox = slide.shapes.add_textbox(Inches(x), Inches(1.55), Inches(4.5), Inches(5.2))
        tf = bbox.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item
            p.font.size = Pt(15)
            p.font.color.rgb = DARK
            p.space_after = Pt(8)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "PageIndex — Project Handoff",
        "NSTSCE / VTTI Document Indexing & RAG Chatbot\nReasoning-based retrieval over research PDFs",
    )

    add_content_slide(
        prs,
        "Project Overview",
        [
            "Modified fork of VectifyAI PageIndex — vectorless, outline-based RAG",
            "Indexes NSTSCE/VTTI research PDFs into hierarchical section trees",
            "Answers natural-language questions with cited sources",
            "Hybrid stack: Ollama (indexing + embeddings) + OpenAI (RAG answers)",
            "Deliverables: CLI tools, batch indexer, two Flask web UIs",
        ],
    )

    add_content_slide(
        prs,
        "What It Does — Two Phases",
        [
            "PHASE 1 — INDEXING (offline): PDF → section tree → summaries → abstract → keywords",
            "  Output: results/*_structure.json + results/DocIndex.json",
            "PHASE 2 — QUERYING (online): question → match docs → tree search → answer",
            "  Output: cited answer in VTTI AI Chatbot (port 5001) or CLI",
            "Unlike vector RAG: navigates document outline with LLM reasoning, not chunk similarity",
        ],
    )

    add_section_slide(prs, "Indexing Pipeline")

    add_content_slide(
        prs,
        "Indexing — 9 Core Steps",
        [
            "1. Parse PDF pages",
            "2. Extract title & authors (rule-based or LLM)",
            "3. Build hierarchical tree (TOC detection, page mapping)",
            "4. Assign node IDs",
            "5. Attach node text (optional)",
            "6–7. Generate leaf & parent summaries",
            "8. Generate document abstract",
            "9. Generate keywords (default: embed_mmr via Ollama)",
            "+ Save structure JSON, merge keywords into DocIndex.json",
        ],
        subtitle="Command: python build_docindex.py  (batch)  |  python run_pageindex.py --pdf_path … --update_docindex",
    )

    add_section_slide(prs, "RAG Query Pipeline")

    add_content_slide(
        prs,
        "RAG — 6 Steps per Question",
        [
            "Step 1: Embed query (+ recent user turns) → match keywords → rank PDFs",
            "Step 2: Load *_structure.json for matched documents",
            "Step 3: LLM tree search — pick relevant section node_ids",
            "Step 4: Extract text/summaries from selected nodes",
            "Step 5: LLM answer per document with [Source: file.pdf] citations",
            "Step 6: Merge multi-doc answers into one response",
            "Web UI sends prior Q&A with each request for follow-up questions",
        ],
        subtitle="Command: python RAG/rag_query.py --query \"…\"  |  python RAG-Interface/app.py",
    )

    add_content_slide(
        prs,
        "Multi-Turn Conversation (Web UI)",
        [
            "Chat UI tracks up to 20 prior turns in the browser session",
            "Each POST /api/query sends: { query, history: [{role, content}, …] }",
            "Step 1: recent user turns combined into retrieval query",
            "Steps 3, 5–6: prior conversation included in LLM prompts",
            "Enables follow-ups like “What about the methodology?” or “Tell me more”",
            "Session-only: refresh clears history; CLI is single-turn by default",
            "Helpers in RAG/utils.py: format_conversation_for_prompt(), build_retrieval_query()",
        ],
    )

    add_content_slide(
        prs,
        "Repository Layout",
        [
            "Database/          — source PDFs (copy files here)",
            "results/           — *_structure.json + DocIndex.json",
            "pageindex/         — core indexing library",
            "build_docindex.py  — batch index all PDFs",
            "run_pageindex.py   — index one PDF or Markdown file",
            "RAG/               — query pipeline (rag_query.py, utils.py)",
            "RAG-Interface/     — VTTI AI Chatbot (port 5001)",
            "DocIndex-Interface/ — doc manager & re-index UI (port 5002)",
            "llm_models.yaml, pageindex/config.yaml, .env — configuration",
        ],
    )

    add_two_column_slide(
        prs,
        "Web Interfaces",
        [
            "VTTI AI Chatbot (port 5001)",
            "python RAG-Interface/app.py",
            "Multi-turn Q&A + session history",
            "SSE progress + source citations",
            "See RAG-Interface/README.md",
        ],
        [
            "DocIndex Manager (port 5002)",
            "python DocIndex-Interface/app.py",
            "Upload PDFs to Database/",
            "Index, re-index, or remove documents",
            "See DocIndex-Interface/README.md",
        ],
        "RAG Chat",
        "DocIndex Manager",
    )

    add_content_slide(
        prs,
        "DocIndex Manager — Document Lifecycle",
        [
            "Upload: drag-and-drop or file picker → saves to Database/",
            "Add Structure: run_pageindex.py --update_docindex for one PDF",
            "Update Structure: re-index after PDF or config changes",
            "Remove: delete Database/ file + structure JSON + DocIndex keywords",
            "Reconstruct: build_docindex.py over all PDFs in Database/",
            "Restart server after code updates (Ctrl+C → python app.py)",
        ],
        subtitle="http://localhost:5002",
    )

    add_section_slide(prs, "Configuration")

    add_content_slide(
        prs,
        "Current Default Setup (Hybrid)",
        [
            "llm_models.yaml:",
            "  indexing.chat_model: qwen7        → Ollama (summaries/abstract)",
            "  indexing.embed_model: mxbai-embed-large",
            "  rag.chat_model: gpt51             → OpenAI (needs CHATGPT_API_KEY)",
            "  rag.tree_search_model: gpt51",
            "  rag.keyword_embed_model: mxbai-embed-large → Ollama",
            "Embeddings ALWAYS use local Ollama — even when RAG chat uses OpenAI",
            "Fully local option: set all rag.* models to qwen7 in llm_models.yaml",
        ],
    )

    add_content_slide(
        prs,
        "Model Routing (Automatic)",
        [
            "qwen7, llama8, ollama …  → Ollama at OLLAMA_BASE_URL",
            "gpt51, gpt4mini, gpt4 …  → OpenAI via CHATGPT_API_KEY in .env",
            "huggingface              → HF_TOKEN in .env",
            "mxbai-embed-large        → Ollama (always)",
            "API_PROVIDER env var is NOT used (legacy docs only)",
            "Restart Flask apps after changing llm_models.yaml",
        ],
    )

    add_section_slide(prs, "Getting Started")

    add_content_slide(
        prs,
        "New Maintainer — 5 Steps",
        [
            "1. pip install -r requirements.txt  (from project root)",
            "2. cp .env.example .env → set CHATGPT_API_KEY (for gpt51 RAG)",
            "3. ollama serve; ollama pull mxbai-embed-large; ollama pull qwen2.5:7b",
            "4. Put PDFs in Database/ → python build_docindex.py",
            "5. python RAG-Interface/app.py  OR  python RAG/rag_query.py --query \"…\"",
            "Always run commands from the PageIndex/ project root",
        ],
    )

    add_content_slide(
        prs,
        "Day-to-Day Operations",
        [
            "Add PDF: upload in DocIndex UI or copy to Database/ → index",
            "Remove PDF: Remove button in DocIndex UI (cleans all artifacts)",
            "Re-index: Update Structure (one file) or DocIndex Reconstruction (all)",
            "Change answer model: edit rag.* in llm_models.yaml → restart chat UI",
            "Customize prompts: RAG/utils.py (answers, tree search, history helpers)",
            "After git pull: restart Flask apps so new API routes load",
            "Logs: logs/ (indexing) | terminal output (web UIs)",
        ],
    )

    add_content_slide(
        prs,
        "Troubleshooting",
        [
            "No DocIndex found → run build_docindex.py from project root",
            "Ollama / embedding errors → ollama serve; pull mxbai-embed-large",
            "OpenAI errors → check CHATGPT_API_KEY for gpt* aliases",
            "DocIndex API 404 / JSON parse error → restart Flask server",
            "No matching documents → re-index; lower keyword-match threshold",
            "Empty doc_authors → re-index PDF (author extraction improved)",
            "GPU driver mismatch → reboot after NVIDIA driver updates",
        ],
    )

    add_content_slide(
        prs,
        "Documentation",
        [
            "README.md — full setup, reference, troubleshooting",
            "RAG-Interface/README.md — chat UI + conversation history",
            "DocIndex-Interface/README.md — upload, index, remove",
            "docs/PageIndex_Handoff.pptx — this slide deck",
            ".env.example — secrets template",
            "scripts/generate_handoff_ppt.py — regenerate deck",
            "benchmarks/ + eval/ — timing and RAGAS evaluation",
        ],
    )

    add_content_slide(
        prs,
        "Handoff Checklist",
        [
            "☐ Repo access transferred",
            "☐ Python 3.12 + pip install -r requirements.txt",
            "☐ .env from .env.example; CHATGPT_API_KEY set if using gpt51",
            "☐ Ollama running; mxbai-embed-large + qwen2.5:7b pulled",
            "☐ build_docindex.py completed; results/DocIndex.json exists",
            "☐ RAG chat: question + follow-up smoke test (port 5001)",
            "☐ DocIndex manager: upload, index, remove smoke test (port 5002)",
            "☐ README.md + handoff PPT reviewed by new maintainer",
        ],
    )

    add_title_slide(prs, "Questions?", "Documentation: README.md  |  Contact: [your name / team]")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
